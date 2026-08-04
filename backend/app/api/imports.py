"""Data migration: bring an existing office's data into the platform.

D-059 F3 turns the plain CSV import into a proper migration wizard — in the
German market switching software costs offices 5–15k € in consultant time,
so this is a sales argument, not a convenience:

- CSV **and XLSX** (RA-MICRO/Advoware export their lists as Excel);
- column-mapping **presets** for the typical RA-MICRO and Advoware exports
  plus a generic one; custom header→field mapping on top;
- `POST /import/preview` — full dry-run (nothing written) that returns the
  detected mapping, normalized sample rows and would-create/skip counts;
- matters: `opposing_party` becomes a real `MatterParty`, the old office's
  file number (Aktenzeichen) lands in `Matter.extra["legacy_reference"]`
  and is searchable in global search;
- `POST /import/documents` — a ZIP with `{reference}/file.pdf` folders is
  matched to matters (by our reference OR the legacy one) and every file
  runs through the normal pipeline (encrypt → OCR → index → deadline scan).

Canonical contacts columns:
    display_name*, kind, email, phone, address, national_id, notes
Canonical matters columns:
    title*, practice_area, status, client_name, opposing_party, court_name,
    court_case_number, opened_on (YYYY-MM-DD), description, legacy_reference
"""
import asyncio
import csv
import io
import json
import logging
import uuid
import zipfile
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.matters import _next_reference
from app.database import SessionLocal, get_db
from app.deps import get_current_user
from app.jurisdictions import jurisdiction_row_for_tenant
from app.models import Contact, Matter, MatterParty, User
from app.models.matters import MATTER_STATUSES, PRACTICE_AREAS
from app.services.audit import log_action

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/import", tags=["import"])

MAX_ROWS = 2000
MAX_ZIP_BYTES = 500 * 1024 * 1024
PREVIEW_SAMPLE = 10

# ---------------------------------------------------------------------------
# Column aliases (all lowercase). Base covers our canonical names + the
# Albanian/German words people actually put in home-grown spreadsheets;
# presets add the header names of the typical RA-MICRO / Advoware exports.
# Everything here is data — adding a legacy system means adding a dict.
# ---------------------------------------------------------------------------

CONTACT_ALIASES: dict[str, list[str]] = {
    "display_name": ["display_name", "emri", "name", "fullname", "full_name"],
    "kind": ["kind", "lloji", "art", "typ"],
    "email": ["email", "e-mail", "e-mail-adresse", "emaili", "mail"],
    "phone": ["phone", "telefoni", "telefon", "telefon1", "telefonnummer", "mobil", "handy"],
    "address": ["address", "adresa", "adresse", "anschrift"],
    "national_id": ["national_id", "nr_personal", "numri_personal", "steuernummer", "geburtsdatum_id"],
    "notes": ["notes", "shenime", "shënime", "bemerkung", "bemerkungen", "notizen", "notiz"],
}

MATTER_ALIASES: dict[str, list[str]] = {
    "title": ["title", "titulli", "aktenkurzbezeichnung", "kurzbezeichnung", "wegen", "betreff", "bezeichnung"],
    "practice_area": ["practice_area", "fusha", "referat", "sachgebiet", "rechtsgebiet"],
    "status": ["status", "statusi", "aktenstatus"],
    "client_name": ["client_name", "klienti", "mandant", "mandantenname", "auftraggeber"],
    "opposing_party": ["opposing_party", "pala_kundershtare", "gegner", "gegnername", "gegenseite"],
    "court_name": ["court_name", "gjykata", "gericht", "instanz"],
    "court_case_number": ["court_case_number", "az_gericht", "gerichtliches_az", "gerichts-az", "az_des_gerichts"],
    "opened_on": ["opened_on", "data_hapjes", "angelegt", "anlagedatum", "datum"],
    "description": ["description", "pershkrimi", "përshkrimi", "beschreibung", "gegenstand"],
    "legacy_reference": ["legacy_reference", "referenca_vjeter", "aktenzeichen", "az", "aktennummer", "akten-nr"],
}

# Preset = extra aliases + composition hints. RA-MICRO "Adressen"/"Akten"
# and Advoware address/file exports carry split name & address columns.
PRESETS: dict[str, dict] = {
    "generic": {"label": "CSV/XLSX (generic)", "contacts": {}, "matters": {}},
    "ra_micro": {
        "label": "RA-MICRO (Adressen/Akten-Export)",
        "contacts": {
            "display_name": ["nachname"],
            "phone": ["telefon geschäftlich", "telefon privat"],
        },
        "matters": {"legacy_reference": ["aktenzeichen"], "title": ["wegen"]},
    },
    "advoware": {
        "label": "Advoware (Adress-/Aktenliste)",
        "contacts": {"display_name": ["nachname/firma", "nachname"]},
        "matters": {"legacy_reference": ["aktenzeichen"], "title": ["kurzrubrum"]},
    },
}

# Split columns that compose a canonical value when it is not mapped directly
_NAME_PARTS = ("vorname", "nachname", "name", "nachname/firma")
_ADDRESS_PARTS = ("straße", "strasse", "plz", "ort")

# Practice-area values seen in DE/XK exports -> our canonical areas
AREA_ALIASES: dict[str, str] = {
    **{a: a for a in PRACTICE_AREAS},
    "civile": "civil", "zivilrecht": "civil", "zivil": "civil",
    "mietrecht": "civil", "verkehrsrecht": "civil", "erbrecht": "civil",
    "penale": "criminal", "strafrecht": "criminal", "strafsachen": "criminal",
    "familjare": "family", "familienrecht": "family", "familiensachen": "family",
    "komerciale": "commercial", "handelsrecht": "commercial",
    "gesellschaftsrecht": "commercial", "wirtschaftsrecht": "commercial",
    "pune": "labor", "punë": "labor", "arbeitsrecht": "labor", "arbeitssachen": "labor",
    "administrative": "administrative", "verwaltungsrecht": "administrative",
    "sozialrecht": "administrative",
}

STATUS_ALIASES: dict[str, str] = {
    **{s: s for s in MATTER_STATUSES},
    "hapur": "open", "offen": "open", "laufend": "open", "aktiv": "open",
    "mbyllur": "closed", "abgeschlossen": "closed", "beendet": "closed",
    "abgelegt": "archived", "archiviert": "archived", "arkivuar": "archived",
}


# ---------------------------------------------------------------------------
# File reading: CSV (utf-8 / latin-1) and XLSX, both into list[dict] with
# lowercase-stripped headers.
# ---------------------------------------------------------------------------

def _read_csv(content: bytes) -> list[dict]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    # RA-MICRO/Advoware CSV exports are usually semicolon-separated
    sample = text[:4096]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    if not reader.fieldnames:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "CSV has no header row")
    return [
        {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
        for row in reader
    ]


def _read_xlsx(content: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(
            status.HTTP_501_NOT_IMPLEMENTED, "XLSX support not installed"
        )
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Unreadable XLSX file")
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "XLSX has no header row")
    keys = [str(h or "").strip().lower() for h in header]

    def _cell(v) -> str:
        if v is None:
            return ""
        if isinstance(v, datetime):
            return v.date().isoformat()
        if isinstance(v, date):
            return v.isoformat()
        return str(v).strip()

    rows = []
    for raw in rows_iter:
        row = {k: _cell(v) for k, v in zip(keys, raw) if k}
        if any(row.values()):
            rows.append(row)
        if len(rows) > MAX_ROWS:
            break
    wb.close()
    return rows


def _read_rows(filename: str, content: bytes) -> list[dict]:
    name = (filename or "").lower()
    rows = _read_xlsx(content) if name.endswith((".xlsx", ".xlsm")) else _read_csv(content)
    if len(rows) > MAX_ROWS:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Too many rows (max {MAX_ROWS})")
    return rows


# ---------------------------------------------------------------------------
# Normalization: source rows -> canonical rows, via preset aliases and an
# optional custom {source_header: canonical_field} mapping from the wizard.
# ---------------------------------------------------------------------------

def _aliases_for(kind: str, preset: str) -> dict[str, list[str]]:
    base = CONTACT_ALIASES if kind == "contacts" else MATTER_ALIASES
    extra = PRESETS.get(preset, PRESETS["generic"]).get(kind, {})
    merged = {field: list(names) for field, names in base.items()}
    for field, names in extra.items():
        # Preset aliases take precedence — put them first
        merged[field] = list(names) + merged.get(field, [])
    return merged


def _auto_map(headers: list[str], kind: str, preset: str) -> dict[str, str]:
    """{source_header: canonical_field} — first alias match wins per field."""
    aliases = _aliases_for(kind, preset)
    mapping: dict[str, str] = {}
    for field, names in aliases.items():
        for name in names:
            if name in headers and name not in mapping:
                mapping[name] = field
                break
    return mapping


def _compose(row: dict, canonical: dict) -> dict:
    """Fill display_name/address from split German columns when unmapped."""
    if not canonical.get("display_name"):
        parts = [row.get(p, "") for p in ("vorname",)] + [
            row.get(p, "") for p in ("nachname", "name", "nachname/firma")
        ]
        joined = " ".join(p for p in parts if p).strip()
        if joined:
            canonical["display_name"] = joined
    if not canonical.get("address"):
        street = row.get("straße") or row.get("strasse") or ""
        town = " ".join(p for p in (row.get("plz", ""), row.get("ort", "")) if p)
        joined = ", ".join(p for p in (street, town) if p)
        if joined:
            canonical["address"] = joined
    return canonical


def _parse_opened(raw: str) -> date | None:
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


def _normalize(
    rows: list[dict], kind: str, preset: str, custom: dict[str, str] | None
) -> tuple[list[dict], dict[str, str], list[str]]:
    """Returns (canonical_rows, effective_mapping, unmapped_headers)."""
    headers = sorted({h for row in rows for h in row})
    mapping = _auto_map(headers, kind, preset)
    for source, field in (custom or {}).items():
        source = source.strip().lower()
        valid = CONTACT_ALIASES if kind == "contacts" else MATTER_ALIASES
        if field in valid:
            mapping[source] = field
        elif not field:  # explicit "ignore this column"
            mapping.pop(source, None)
    normalized = []
    for row in rows:
        canonical: dict[str, str] = {}
        for source, field in mapping.items():
            value = row.get(source, "")
            if value and not canonical.get(field):
                canonical[field] = value
        if kind == "contacts":
            canonical = _compose(row, canonical)
        normalized.append(canonical)
    unmapped = [h for h in headers if h not in mapping]
    return normalized, mapping, unmapped


# ---------------------------------------------------------------------------
# The import runs — one code path for preview (dry run) and execute, so the
# preview numbers are exactly what the real run will do.
# ---------------------------------------------------------------------------

async def _run_contacts(
    rows: list[dict], user: User, db: AsyncSession, dry_run: bool
) -> dict:
    existing = {
        name.lower()
        for (name,) in (
            await db.execute(
                select(Contact.display_name).where(Contact.tenant_id == user.tenant_id)
            )
        ).all()
    }
    created, skipped, errors = 0, 0, []
    for i, row in enumerate(rows, start=2):  # 1 = header line
        name = row.get("display_name") or ""
        if not name:
            errors.append(f"Line {i}: missing display_name")
            continue
        if name.lower() in existing:
            skipped += 1
            continue
        if not dry_run:
            db.add(
                Contact(
                    tenant_id=user.tenant_id,
                    display_name=name[:255],
                    kind=(row.get("kind") or "person")[:16],
                    email=(row.get("email") or None),
                    phone=(row.get("phone") or None),
                    address=(row.get("address") or None),
                    national_id=(row.get("national_id") or None),
                    notes=(row.get("notes") or None),
                )
            )
        existing.add(name.lower())
        created += 1
    return {"created": created, "skipped": skipped, "errors": errors}


async def _run_matters(
    rows: list[dict], user: User, db: AsyncSession, dry_run: bool
) -> dict:
    jurisdiction = await jurisdiction_row_for_tenant(db, user.tenant_id)
    contact_names = {
        name.lower()
        for (name,) in (
            await db.execute(
                select(Contact.display_name).where(Contact.tenant_id == user.tenant_id)
            )
        ).all()
    }

    async def _find_or_create_contact(name: str) -> uuid.UUID | None:
        if dry_run:
            contact_names.add(name.lower())
            return None
        contact = (
            await db.execute(
                select(Contact).where(
                    Contact.tenant_id == user.tenant_id,
                    func.lower(Contact.display_name) == name.lower(),
                )
            )
        ).scalars().first()
        if contact is None:
            contact = Contact(tenant_id=user.tenant_id, display_name=name[:255])
            db.add(contact)
            await db.flush()
        contact_names.add(name.lower())
        return contact.id

    created, errors, new_contacts = 0, [], 0
    for i, row in enumerate(rows, start=2):
        title = row.get("title") or row.get("legacy_reference") or ""
        if not title:
            errors.append(f"Line {i}: missing title")
            continue
        area = AREA_ALIASES.get((row.get("practice_area") or "civil").lower(), "other")
        matter_status = STATUS_ALIASES.get((row.get("status") or "open").lower(), "open")
        opened_on: date | None = None
        raw_date = row.get("opened_on") or ""
        if raw_date:
            opened_on = _parse_opened(raw_date)
            if opened_on is None:
                errors.append(f"Line {i}: unreadable date '{raw_date}' (use YYYY-MM-DD)")

        client_name = row.get("client_name") or ""
        opposing_name = row.get("opposing_party") or ""
        for n in (client_name, opposing_name):
            if n and n.lower() not in contact_names:
                new_contacts += 1
        client_id = await _find_or_create_contact(client_name) if client_name else None
        opposing_id = (
            await _find_or_create_contact(opposing_name) if opposing_name else None
        )

        if not dry_run:
            extra = {}
            if row.get("legacy_reference"):
                # The old office's Aktenzeichen — searchable in global search
                extra["legacy_reference"] = row["legacy_reference"][:64]
            matter = Matter(
                tenant_id=user.tenant_id,
                reference=await _next_reference(db, user.tenant_id, area),
                title=title[:500],
                description=(row.get("description") or None),
                practice_area=area,
                status=matter_status,
                jurisdiction_id=jurisdiction.id,
                responsible_user_id=user.id,
                client_id=client_id,
                court_name=(row.get("court_name") or None),
                court_case_number=(row.get("court_case_number") or None),
                opened_on=opened_on or date.today(),
                extra=extra,
            )
            db.add(matter)
            await db.flush()
            if opposing_id:
                db.add(
                    MatterParty(
                        matter_id=matter.id,
                        contact_id=opposing_id,
                        role="opposing_party",
                    )
                )
        created += 1
    return {"created": created, "errors": errors, "new_contacts": new_contacts}


def _parse_mapping(mapping: str | None) -> dict[str, str] | None:
    if not mapping:
        return None
    try:
        parsed = json.loads(mapping)
        if not isinstance(parsed, dict):
            raise ValueError
        return {str(k): str(v) for k, v in parsed.items()}
    except (json.JSONDecodeError, ValueError):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "mapping must be a JSON object")


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

@router.get("/presets")
async def list_presets(user: User = Depends(get_current_user)):
    """The wizard's source step: available presets + canonical fields."""
    return {
        "presets": [
            {"id": pid, "label": p["label"]} for pid, p in PRESETS.items()
        ],
        "fields": {"contacts": list(CONTACT_ALIASES), "matters": list(MATTER_ALIASES)},
    }


@router.post("/preview")
async def preview_import(
    kind: str = Form(...),
    preset: str = Form("generic"),
    mapping: str | None = Form(None),
    file: UploadFile = File(...),
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Dry run: parse, map and validate WITHOUT writing anything. The wizard
    shows the mapping and the would-create/skip counts before executing."""
    if kind not in ("contacts", "matters"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "kind must be contacts|matters")
    rows = _read_rows(file.filename or "", await file.read())
    normalized, eff_mapping, unmapped = _normalize(
        rows, kind, preset, _parse_mapping(mapping)
    )
    runner = _run_contacts if kind == "contacts" else _run_matters
    result = await runner(normalized, user, db, dry_run=True)
    await db.rollback()  # belt & braces — the dry run must leave no trace
    return {
        "rows_total": len(rows),
        "mapping": eff_mapping,
        "unmapped": unmapped,
        "sample": normalized[:PREVIEW_SAMPLE],
        **{k: v for k, v in result.items() if k != "errors"},
        "errors": result["errors"][:20],
    }


@router.post("/contacts")
async def import_contacts(
    file: UploadFile = File(...),
    preset: str = Form("generic"),
    mapping: str | None = Form(None),
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    rows = _read_rows(file.filename or "", await file.read())
    normalized, _, _ = _normalize(rows, "contacts", preset, _parse_mapping(mapping))
    result = await _run_contacts(normalized, user, db, dry_run=False)
    await log_action(
        db, tenant_id=user.tenant_id, user_id=user.id, action="import.contacts",
        entity_type="tenant", entity_id=str(user.tenant_id),
        details={"created": result["created"], "skipped": result["skipped"], "preset": preset},
    )
    await db.commit()
    return {**result, "errors": result["errors"][:20]}


@router.post("/matters")
async def import_matters(
    file: UploadFile = File(...),
    preset: str = Form("generic"),
    mapping: str | None = Form(None),
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    rows = _read_rows(file.filename or "", await file.read())
    normalized, _, _ = _normalize(rows, "matters", preset, _parse_mapping(mapping))
    result = await _run_matters(normalized, user, db, dry_run=False)
    await log_action(
        db, tenant_id=user.tenant_id, user_id=user.id, action="import.matters",
        entity_type="tenant", entity_id=str(user.tenant_id),
        details={"created": result["created"], "preset": preset},
    )
    await db.commit()
    return {**result, "errors": result["errors"][:20]}


# ---------------------------------------------------------------------------
# Bulk documents: ZIP with one folder per matter.
# ---------------------------------------------------------------------------

async def _process_document_zip(
    zip_bytes: bytes, matched: dict[str, uuid.UUID], tenant_id: uuid.UUID, user_id: uuid.UUID
) -> None:
    """Background: run every matched file through the normal document
    pipeline (encrypt → OCR/text → index → deadline scan). Own DB session —
    the request that scheduled us is long gone."""
    from app.api.documents import upload_type_allowed
    from app.models import Document
    from app.services import deadline_suggest, rag, storage

    done, failed = 0, 0
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            parts = info.filename.split("/")
            if len(parts) < 2 or parts[0] not in matched:
                continue
            filename = parts[-1]
            if filename.startswith(".") or not upload_type_allowed(filename, None):
                continue
            content = zf.read(info)
            if not content:
                continue
            matter_id = matched[parts[0]]
            try:
                async with SessionLocal() as db:
                    rel_path, sha, encrypted = storage.save_file(content, tenant_id)
                    mime = "application/octet-stream"
                    text, ocr_status = storage.extract_text(content, mime, filename)
                    doc = Document(
                        tenant_id=tenant_id,
                        matter_id=matter_id,
                        family_id=uuid.uuid4(),
                        version=1,
                        is_latest=True,
                        filename=filename,
                        mime_type=mime,
                        size_bytes=len(content),
                        sha256=sha,
                        storage_path=rel_path,
                        encrypted=encrypted,
                        category="import",
                        uploaded_by_id=user_id,
                        text_content=text or None,
                        ocr_status=ocr_status,
                    )
                    db.add(doc)
                    await db.flush()
                    if text:
                        await rag.index_document_chunks(db, doc)
                    await db.commit()
                    doc_id = doc.id
                if text:
                    deadline_suggest.scan_document_background(doc_id)
                done += 1
            except Exception as exc:  # noqa: BLE001 — keep the batch going
                logger.warning("Document import failed for %s: %s", info.filename, exc)
                failed += 1
            await asyncio.sleep(0.05)  # don't starve the event loop (OCR is heavy)

    async with SessionLocal() as db:
        await log_action(
            db, tenant_id=tenant_id, user_id=user_id, action="import.documents",
            entity_type="tenant", entity_id=str(tenant_id),
            details={"imported": done, "failed": failed, "folders": len(matched)},
        )
        await db.commit()
    logger.info("Document ZIP import done: %s ok, %s failed", done, failed)


@router.post("/documents")
async def import_documents(
    file: UploadFile = File(...),
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """ZIP upload: top-level folders are matched to matters by our reference
    or the imported legacy_reference; files are processed in the background
    through the standard document pipeline. Returns the match summary
    immediately — processing continues server-side."""
    content = await file.read()
    if len(content) > MAX_ZIP_BYTES:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "ZIP too large")
    try:
        zf = zipfile.ZipFile(io.BytesIO(content))
        names = zf.namelist()
    except zipfile.BadZipFile:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Not a valid ZIP file")

    folders = {n.split("/")[0] for n in names if "/" in n}
    if not folders:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "ZIP must contain one folder per matter: {reference}/file.pdf",
        )

    matters = (
        await db.execute(select(Matter).where(Matter.tenant_id == user.tenant_id))
    ).scalars().all()
    by_ref: dict[str, uuid.UUID] = {}
    for m in matters:
        by_ref[m.reference.lower()] = m.id
        legacy = (m.extra or {}).get("legacy_reference")
        if legacy:
            by_ref[str(legacy).lower()] = m.id
    # Filenames use "_" where references have "/" (FAT/zip-unsafe character)
    matched = {
        folder: by_ref[key]
        for folder in folders
        if (key := folder.lower().replace("_", "/")) in by_ref
        or (key := folder.lower()) in by_ref
    }
    unmatched = sorted(folders - set(matched))
    file_count = sum(
        1 for n in names
        if "/" in n and n.split("/")[0] in matched and not n.endswith("/")
    )

    if matched:
        asyncio.create_task(
            _process_document_zip(content, matched, user.tenant_id, user.id)
        )
    return {
        "matched_folders": len(matched),
        "unmatched_folders": unmatched[:50],
        "files_queued": file_count,
    }
